"""Security/behavior regression checks with generated, non-private fixtures."""
import hashlib
import io
import json
from pathlib import Path
import tempfile
import unittest

import pymupdf as fitz
import pikepdf
from PIL import Image
from engine import run, PdfError


class EngineTests(unittest.TestCase):
    def test_cancellation_still_stops_owned_worker_if_taskkill_times_out(self):
        import os, subprocess, sys
        from unittest.mock import patch
        import engine
        if os.name != 'nt':
            self.skipTest('Windows helper fallback')
        child=subprocess.Popen([sys.executable,'-c','import time; time.sleep(30)'],creationflags=subprocess.CREATE_NO_WINDOW)
        try:
            with patch.object(engine.subprocess,'run',side_effect=subprocess.TimeoutExpired('taskkill',10)):
                engine.stop_process_tree(child)
            self.assertIsNotNone(child.poll())
        finally:
            if child.poll() is None: child.kill()
            child.wait(timeout=5)

    @classmethod
    def setUpClass(cls):
        cls.root = Path(tempfile.mkdtemp(prefix="pdf-engine-test-"))
        cls.input = cls.root / "Unicode 한글 sample.pdf"
        with fitz.open() as doc:
            for i in range(3):
                page = doc.new_page(width=400, height=500)
                page.insert_text((40, 70), f"PAGE {i + 1} PUBLIC TEXT", fontsize=18)
                page.insert_text((40, 130), "SECRET-ACCOUNT-12345", fontsize=18)
                page.draw_rect((40, 180, 200, 230), color=(1, 0, 0), fill=(1, 0, 0))
            page = doc[0]
            widget = fitz.Widget()
            widget.field_name = "student_name"
            widget.field_type = fitz.PDF_WIDGET_TYPE_TEXT
            widget.rect = fitz.Rect(40, 280, 270, 310)
            widget.field_value = "OLD NAME"
            page.add_widget(widget)
            doc.set_metadata({"title": "SECRET-ACCOUNT-12345"})
            doc.embfile_add("private.txt", b"SECRET-ACCOUNT-12345")
            doc.save(cls.input)
        cls.source_hash = hashlib.sha256(cls.input.read_bytes()).hexdigest()
        cls.signature = cls.root / "signature.png"
        image = Image.new("RGBA", (120, 50), (0, 0, 255, 255))
        image.save(cls.signature)
        cls.scanned = cls.root / "scan.pdf"
        with fitz.open() as source, fitz.open() as target:
            page = source.new_page(width=600, height=240)
            page.insert_text((30, 90), "APPLE ENGLISH 123", fontsize=34)
            pix = page.get_pixmap(dpi=150)
            dst = target.new_page(width=600, height=240)
            dst.insert_image(dst.rect, pixmap=pix)
            target.save(cls.scanned)

    def invoke(self, action, **options):
        return run(action, {"inputs": [str(self.input)], "outputDir": str(self.root), **options})

    def path(self, result, index=0):
        self.assertTrue(result["ok"])
        return Path(result["outputs"][index]["path"])

    def test_01_inspect_preview_and_unicode_path(self):
        info = self.invoke("inspect")["details"]
        self.assertEqual(info["pageCount"], 3)
        self.assertEqual(info["fields"][0]["name"], "student_name")
        preview = self.invoke("preview")["details"]
        self.assertEqual((preview["width"], preview["height"]), (400, 500))
        self.assertTrue(preview["pngBase64"].startswith("iVBOR"))

    def test_02_protect_wrong_password_unlock_and_permissions(self):
        secured = self.path(self.invoke("protect", newPassword="correct-password", permissions={"copy": False}))
        self.assertTrue(self.invoke("inspect", inputs=[str(secured)])["details"]["needsPassword"])
        with self.assertRaises(PdfError) as ctx:
            self.invoke("unlock", inputs=[str(secured)], password="incorrect")
        self.assertEqual(ctx.exception.code, "PASSWORD_REQUIRED")
        unlocked = self.path(self.invoke("unlock", inputs=[str(secured)], password="correct-password"))
        with fitz.open(unlocked) as doc:
            self.assertFalse(doc.is_encrypted)
            self.assertIn("PUBLIC TEXT", doc[0].get_text())
        ownerlocked = self.root / "owner-only.pdf"
        with pikepdf.open(self.input) as doc:
            doc.save(ownerlocked, encryption=pikepdf.Encryption(user="", owner="owner-password", R=6))
        self.path(self.invoke("unlock", inputs=[str(ownerlocked)]))

    def test_03_page_operations(self):
        reordered = self.path(self.invoke("reorder", pages="3,1,1,2"))
        with fitz.open(reordered) as doc:
            self.assertEqual(doc.page_count, 4)
            self.assertIn("PAGE 3", doc[0].get_text())
            self.assertIn("PAGE 1", doc[2].get_text())
        self.assertEqual(len(self.invoke("split", every=2)["outputs"]), 2)
        extracted = self.path(self.invoke("extract", pages="2-3"))
        with fitz.open(extracted) as doc:
            self.assertEqual(doc.page_count, 2)
            self.assertIn("PAGE 2", doc[0].get_text())
        merged = self.path(self.invoke("merge", inputs=[str(self.input), str(extracted)]))
        with fitz.open(merged) as doc:
            self.assertEqual(doc.page_count, 5)
        rotated = self.path(self.invoke("rotate", pages="2", angle=90))
        with fitz.open(rotated) as doc:
            self.assertEqual(doc[0].rotation, 0)
            self.assertEqual(doc[1].rotation, 90)
        self.path(self.invoke("compress"))
        with self.assertRaises(PdfError):
            self.invoke("extract", pages="0,999")

    def test_04_image_conversions(self):
        rendered = self.invoke("pdf-to-images", pages="1", dpi=72)
        image = self.path(rendered)
        with Image.open(image) as pic:
            self.assertEqual(pic.size, (400, 500))
        pdf = self.path(self.invoke("images-to-pdf", inputs=[str(image), str(self.signature)], pageSize="A4"))
        with fitz.open(pdf) as doc:
            self.assertEqual(doc.page_count, 2)
            self.assertAlmostEqual(doc[0].rect.width, 595.276, places=2)

    def test_05_overlays_and_forms(self):
        watermark = self.path(self.invoke("watermark", text="INTERNAL REVIEW"))
        with fitz.open(watermark) as doc:
            self.assertIn("INTERNAL REVIEW", doc[0].get_text())
        numbered = self.path(self.invoke("page-numbers", start=5))
        with fitz.open(numbered) as doc:
            self.assertIn("5 / 3", doc[0].get_text())
        filled = self.path(self.invoke("fill-form", values={"student_name": "Alice"}))
        with fitz.open(filled) as doc:
            self.assertEqual(next(doc[0].widgets()).field_value, "Alice")
        flat = self.path(self.invoke("fill-form", values={"student_name": "Alice"}, flatten=True))
        with fitz.open(flat) as doc:
            self.assertEqual(list(doc[0].widgets() or []), [])
            self.assertIn("Alice", doc[0].get_text())
        signed = self.path(self.invoke("sign", signaturePath=str(self.signature), page=1, box={"x0": 40, "y0": 330, "x1": 220, "y1": 405}))
        with fitz.open(signed) as doc:
            self.assertTrue(doc[0].get_images())

    def test_06_true_redaction_and_document_scrub(self):
        with self.assertRaises(PdfError):
            self.invoke("redact", searchText="SECRET-ACCOUNT-12345")
        for flatten in (False, True):
            redacted = self.path(self.invoke("redact", searchText="SECRET-ACCOUNT-12345", confirmed=True, rasterize=flatten))
            with fitz.open(redacted) as doc:
                self.assertEqual(doc.embfile_count(), 0)
                self.assertNotIn("SECRET-ACCOUNT-12345", json.dumps(doc.metadata))
                for page in doc:
                    self.assertNotIn("SECRET-ACCOUNT-12345", page.get_text())
                if flatten:
                    self.assertEqual(doc[0].get_text(), "")
                else:
                    self.assertIn("PUBLIC TEXT", doc[0].get_text())
            with pikepdf.open(redacted) as doc:
                for obj in doc.objects:
                    if isinstance(obj, pikepdf.Stream):
                        try:
                            self.assertNotIn(b"SECRET-ACCOUNT-12345", obj.read_bytes())
                        except pikepdf.PdfError:
                            pass
        box = self.path(self.invoke("redact", boxes=[{"page": 1, "x0": 40, "y0": 180, "x1": 200, "y1": 230}], confirmed=True))
        with fitz.open(box) as doc:
            pixel = doc[0].get_pixmap().pixel(100, 200)
            self.assertEqual(pixel[:3], (0, 0, 0))

    def test_07_rotated_coordinates(self):
        rotated = self.path(self.invoke("rotate", pages="1", angle=90))
        preview = self.invoke("preview", inputs=[str(rotated)])["details"]
        self.assertEqual((preview["width"], preview["height"]), (500, 400))
        redacted = self.path(self.invoke("redact", inputs=[str(rotated)], boxes=[{"page": 1, "x0": 50, "y0": 50, "x1": 100, "y1": 100}], confirmed=True))
        with fitz.open(redacted) as doc:
            self.assertEqual(doc[0].get_pixmap().pixel(75, 75)[:3], (0, 0, 0))
        self.path(self.invoke("page-numbers", inputs=[str(rotated)]))
        self.path(self.invoke("sign", inputs=[str(rotated)], signaturePath=str(self.signature), page=1, box={"x0": 300, "y0": 200, "x1": 450, "y1": 250}))

    def test_08_offline_ocr(self):
        result = self.invoke("ocr", inputs=[str(self.scanned)], dpi=150)
        output = self.path(result)
        self.assertGreater(result["details"]["recognizedLines"], 0)
        with fitz.open(output) as doc:
            self.assertIn("APPLE", doc[0].get_text().upper())
        with fitz.open(output) as out, fitz.open(self.scanned) as original:
            self.assertEqual(out[0].get_pixmap().samples, original[0].get_pixmap().samples)

    def test_09_inputs_preserved_outputs_unique_validation(self):
        one = self.path(self.invoke("extract-text"))
        two = self.path(self.invoke("extract-text"))
        self.assertNotEqual(one, two)
        self.assertEqual(hashlib.sha256(self.input.read_bytes()).hexdigest(), self.source_hash)
        with self.assertRaises(PdfError):
            self.invoke("extract", inputs=["relative.pdf"], pages="1")
        with self.assertRaises(PdfError):
            self.invoke("sign", signaturePath=str(self.signature), box={"x0": -1, "y0": 0, "x1": 100, "y1": 100})
        with self.assertRaises(PdfError):
            self.invoke("shell")

    def test_10_extended_page_and_edit_tools(self):
        removed = self.path(self.invoke("remove-pages", pages="2"))
        with fitz.open(removed) as doc:
            self.assertEqual(doc.page_count, 2)
            self.assertIn("PAGE 3", doc[1].get_text())
        with self.assertRaises(PdfError):
            self.invoke("remove-pages", pages="1-3")
        self.path(self.invoke("repair"))
        cropped = self.path(self.invoke("crop", pages="1", box={"x0": 20, "y0": 20, "x1": 380, "y1": 480}))
        with fitz.open(cropped) as doc:
            self.assertEqual(tuple(doc[0].rect), (0, 0, 360, 460))
        added = self.path(self.invoke("add-text", page=1, text="ADDED CONTENT", box={"x0": 40, "y0": 380, "x1": 350, "y1": 450}))
        with fitz.open(added) as doc:
            self.assertIn("ADDED CONTENT", doc[0].get_text())
        annotated = self.path(self.invoke("annotate", page=1, text="Review note", box={"x0": 40, "y0": 380, "x1": 350, "y1": 450}, kind="note"))
        with fitz.open(annotated) as doc:
            self.assertEqual(next(doc[0].annots()).info["content"], "Review note")
        self.path(self.invoke("scan-to-pdf", inputs=[str(self.signature)], grayscale=True))

    def test_11_document_conversions(self):
        word = self.path(self.invoke("pdf-to-docx", pages="1"))
        from docx import Document
        document = Document(word)
        text = "\n".join(p.text for p in document.paragraphs)
        text += "\n".join(cell.text for table in document.tables for row in table.rows for cell in row.cells)
        self.assertIn("PUBLIC TEXT", text)
        slides = self.path(self.invoke("pdf-to-pptx", pages="1-2", dpi=72))
        from pptx import Presentation
        presentation = Presentation(slides)
        self.assertEqual(len(presentation.slides), 2)
        self.assertEqual(len(presentation.slides[0].shapes), 1)
        workbook_path = self.path(self.invoke("pdf-to-xlsx", pages="1"))
        from openpyxl import load_workbook
        workbook = load_workbook(workbook_path)
        self.assertTrue(workbook.worksheets)
        for sheet in workbook:
            for row in sheet:
                self.assertFalse(any(cell.data_type == "f" for cell in row))
        structured = self.path(self.invoke("extract-structured", format="json"))
        self.assertEqual(json.loads(structured.read_text(encoding="utf-8"))["pageCount"], 3)
        markdown = self.path(self.invoke("extract-structured", format="markdown"))
        self.assertIn("## Page 1", markdown.read_text(encoding="utf-8"))
        self.path(self.invoke("extract-images", inputs=[str(self.scanned)]))

    def test_12_compare_summary_html_and_optional_capabilities(self):
        result = self.invoke("compare", inputs=[str(self.input), str(self.input)])
        self.assertFalse(result["details"]["textChanged"])
        self.assertEqual(result["details"]["visuallyChangedPages"], [])
        self.path(self.invoke("summarize", sentences=2))
        html = self.root / "local.html"
        html.write_text("<html><body><h1>Local document</h1><p>Readable text.</p></body></html>", encoding="utf-8")
        rendered = self.path(self.invoke("html-to-pdf", inputs=[str(html)]))
        with fitz.open(rendered) as doc:
            self.assertIn("Local document", doc[0].get_text())
        html.write_text('<img src="https://example.com/private">', encoding="utf-8")
        with self.assertRaises(PdfError):
            self.invoke("html-to-pdf", inputs=[str(html)])
        # Full model/conformance/signature execution has dedicated integration
        # suites. Keep the basic engine regression focused on readiness routing.
        capabilities = run("capabilities", {})["details"]
        for action in ("pdfa", "translate", "ai-extract", "certificate-sign"):
            operation = next(item for item in capabilities["operations"] if item["id"] == action)
            self.assertEqual(operation["available"], action in capabilities["actions"])

    def test_13_bounded_candidate_testing_and_merge_passwords(self):
        secured = self.path(self.invoke("protect", newPassword="candidate-hit"))
        result = self.invoke("recover-password", inputs=[str(secured)], candidates=["incorrect", "candidate-hit"], maxSeconds=10)
        self.assertEqual(result["details"]["matchedCandidate"], 2)
        self.assertNotIn("candidate-hit", json.dumps(result))
        self.assertEqual(result["details"]["attempts"], 2)
        with self.assertRaises(PdfError) as ctx:
            self.invoke("recover-password", inputs=[str(secured)], candidates=["incorrect"], maxSeconds=1)
        self.assertEqual(ctx.exception.code, "NO_PASSWORD_MATCH")
        merged = self.path(self.invoke("merge", inputs=[str(secured), str(self.input)], passwords=["candidate-hit", ""]))
        with fitz.open(merged) as doc:
            self.assertEqual(doc.page_count, 6)

    def test_14_bundled_office_to_pdf(self):
        self.assertTrue(run("capabilities", {})["details"]["office"]["available"])
        from docx import Document
        source = self.root / "office conversion 한글.docx"
        document = Document()
        document.add_heading("OFFICE CONVERSION TEST", 0)
        document.add_paragraph("The full local document is rendered by LibreOffice.")
        document.save(source)
        original = hashlib.sha256(source.read_bytes()).hexdigest()
        result = self.invoke("office-to-pdf", inputs=[str(source)])
        target = self.path(result)
        with fitz.open(target) as doc:
            self.assertIn("OFFICE CONVERSION TEST", " ".join(doc[0].get_text().split()))
        self.assertEqual(hashlib.sha256(source.read_bytes()).hexdigest(), original)
        self.assertFalse((target.parent / "office-profile").exists())


if __name__ == "__main__":
    unittest.main(verbosity=2)
